#!/usr/bin/env python3
"""
cluster_energy_billing.py
=========================

Monthly energy billing for a shared solar + battery microgrid ("cluster") of
households behind one supplier (retailer) connection.

Pipeline
  1. Households table       -> unit entitlement (UE) shares, battery allocation
  2. Meter export CSVs      -> cumulative register reads -> monthly kWh per meter
  3. Supplier PDF invoices  -> "Usage Calculations" section parsed automatically
  4. Monthly allocation     -> microgrid / grid split of every household's import and
                               export, shared supplier costs, common loads, reconciliation
  5. Exports                -> monthly tabs (.xlsx), one A4 PDF invoice per household,
                               Xero import CSVs, audit CSVs

Example (billing Apr-Jun: also pass the MARCH csv, which provides the opening reads)
  python3 cluster_energy_billing.py --year 2026 --months Apr-Jun \\
      --meters meters_2026-03.csv meters_2026-04.csv meters_2026-05.csv meters_2026-06.csv \\
      --pdfs supplier_2026-04.pdf supplier_2026-05.pdf supplier_2026-06.pdf

Why the extra month: the meters report cumulative registers, so a month's usage is
its closing read minus the previous month's closing read. Always include the CSV of
the month before your first target month. (Alternatively pass a previous run's
<outdir>/meter_readings_ledger.csv with --opening-readings.)

  python3 cluster_energy_billing.py --help    lists every option.

Requirements: Python 3.9+, numpy, pandas, openpyxl, reportlab, and pdfplumber (only if
poppler's `pdftotext` is not installed). Missing packages are installed automatically
on first run; use --no-auto-install to turn this off. See README.md.
"""
from __future__ import annotations

import argparse
import datetime as dt
import os
import re
import shutil
import subprocess
import sys
import unicodedata
from dataclasses import dataclass, field


# -----------------------------------------------------------------------------
# Dependency check / auto-install (runs before the third-party imports below)
# -----------------------------------------------------------------------------
def required_packages() -> dict[str, str]:
    """import name -> pip name. pdfplumber is only needed when poppler's pdftotext is absent."""
    req = {"numpy": "numpy", "pandas": "pandas", "openpyxl": "openpyxl", "reportlab": "reportlab"}
    if shutil.which("pdftotext") is None:
        req["pdfplumber"] = "pdfplumber"
    return req


VENV_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".billing-venv")


def _venv_python(venv: str) -> str:
    return os.path.join(venv, "Scripts", "python.exe") if os.name == "nt" else os.path.join(venv, "bin", "python")


def ensure_dependencies(argv: list[str]) -> None:
    """Make sure the third-party packages are importable.

    * all present             -> nothing happens
    * inside a virtual env     -> pip-install the missing ones there
    * otherwise                -> create a private env (.billing-venv beside this script, reusing any
                                  packages already installed system-wide), install there, and re-run
                                  this script inside it. Avoids "externally-managed-environment" errors
                                  on Homebrew / Debian / Ubuntu Pythons. Later runs reuse the env.
    """
    import importlib
    import importlib.util

    def missing():
        return [pip for mod, pip in required_packages().items() if importlib.util.find_spec(mod) is None]

    todo = missing()
    if not todo:
        return
    manual = ("Install them manually, e.g.:\n  python3 -m venv .venv && . .venv/bin/activate && "
              f"pip install {' '.join(todo)}\nthen run the script with that environment's python.")
    if "--no-auto-install" in argv:
        sys.exit(f"ERROR: missing Python packages: {', '.join(todo)}\n{manual}")

    def pip_install(python: str) -> subprocess.CompletedProcess:
        subprocess.run([python, "-m", "ensurepip", "--upgrade"], capture_output=True)   # no-op if pip exists
        return subprocess.run([python, "-m", "pip", "install", "--disable-pip-version-check", "--quiet", *todo],
                              capture_output=True, text=True)

    def fail(res, where):
        tail = (res.stderr or res.stdout or "").strip().splitlines()
        sys.exit(f"ERROR: could not install {', '.join(todo)} into {where}"
                 f"{': ' + tail[-1] if tail else ''}\n(an internet connection is needed the first time)\n{manual}")

    in_venv = sys.prefix != getattr(sys, "base_prefix", sys.prefix)
    if in_venv:
        print(f"Installing missing Python packages into the active environment: {', '.join(todo)} ...", file=sys.stderr)
        res = pip_install(sys.executable)
        importlib.invalidate_caches()
        if res.returncode != 0 or missing():
            fail(res, sys.prefix)
        print("Packages installed.", file=sys.stderr)
        return

    if os.environ.get("CLUSTER_BILLING_BOOTSTRAPPED"):        # safety net against a re-exec loop
        sys.exit(f"ERROR: packages still missing after setting up {VENV_DIR}: {', '.join(todo)}\n{manual}")
    vpy = _venv_python(VENV_DIR)
    if not os.path.exists(vpy):
        print(f"Missing Python packages: {', '.join(todo)}. Creating a private environment in {VENV_DIR} "
              "(one-off) ...", file=sys.stderr)
        res = subprocess.run([sys.executable, "-m", "venv", "--system-site-packages", VENV_DIR],
                             capture_output=True, text=True)
        if res.returncode != 0:
            fail(res, VENV_DIR + " (could not create it; on Debian/Ubuntu install the 'python3-venv' package)")
    # install whatever the private env still lacks, then re-run the script inside it
    check = subprocess.run([vpy, "-c", "import importlib.util,sys; sys.exit(any(importlib.util.find_spec(m) is None "
                                       f"for m in {list(required_packages())!r}))"])
    if check.returncode != 0:
        print(f"Installing {', '.join(todo)} into {VENV_DIR} ...", file=sys.stderr)
        res = pip_install(vpy)
        if res.returncode != 0:
            fail(res, VENV_DIR)
        print("Packages installed.", file=sys.stderr)
    env = dict(os.environ, CLUSTER_BILLING_BOOTSTRAPPED="1")
    sys.stdout.flush()
    sys.stderr.flush()
    sys.exit(subprocess.run([vpy, os.path.abspath(__file__), *argv], env=env).returncode)


if __name__ == "__main__":
    ensure_dependencies(sys.argv[1:])

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

# -----------------------------------------------------------------------------
# Logging
# -----------------------------------------------------------------------------
NOTES: list[str] = []


def note(msg: str) -> None:
    if msg not in NOTES:
        NOTES.append(msg)
        print(f"NOTE: {msg}", file=sys.stderr)


def info(msg: str) -> None:
    print(msg, file=sys.stderr)


class BillingError(Exception):
    pass


# -----------------------------------------------------------------------------
# Command line
# -----------------------------------------------------------------------------
def build_parser() -> argparse.ArgumentParser:
    here = os.path.dirname(os.path.abspath(__file__))
    p = argparse.ArgumentParser(
        description="Cluster energy billing: monthly tabs + per-household invoices.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter)
    g = p.add_argument_group("Global parameters (tariffs and cluster settings)")
    g.add_argument("--grid-export-rate", type=float, default=0.0385, help="Grid export credit, $/kWh inc GST")
    g.add_argument("--grid-import-rate", type=float, default=0.2994, help="Grid import rate, $/kWh inc GST")
    g.add_argument("--p2p-rate", type=float, default=0.20, help="Microgrid peer-to-peer rate, $/kWh")
    g.add_argument("--num-households", type=int, default=24, help="Households sharing supplier fixed costs")
    g.add_argument("--no-pv-p2p-rate", type=float, default=None,
                   help="No-PV P2P rate. Accepted for compatibility with the original spreadsheet; not used in any calculation")
    g.add_argument("--battery-capacity", type=float, default=232, help="Shared battery (BESS) capacity, kWh")
    g.add_argument("--export-limit", type=float, default=0.66568, help="Export limit constraint factor")
    g.add_argument("--gst", type=float, default=0.10, help="GST rate applied to supplier charges (0.10 = 10%%)")

    r = p.add_argument_group("Run inputs")
    r.add_argument("--year", type=int, default=2026, help="Calendar year of the billing months")
    r.add_argument("--months", required=True, help='Target billing months, e.g. "Apr,May,Jun", "4-6", "Apr-Jun"')
    r.add_argument("--meters", nargs="+", required=True, metavar="CSV",
                   help="Stormcloud meter CSVs (cumulative reads at period end): one per target month PLUS the "
                        "month before the first target month (up to 13)")
    r.add_argument("--pdfs", nargs="+", required=True, metavar="PDF", help="Supplier invoice PDFs, one per month")
    r.add_argument("--opening-readings", nargs="+", default=[], metavar="FILE",
                   help="Optional alternative to the extra CSV: a previous run's meter_readings_ledger.csv")
    r.add_argument("--households", default=os.path.join(here, "households.csv"),
                   help="Households CSV")

    o = p.add_argument_group("Behaviour")
    o.add_argument("--extra-assets", choices=["common", "ignore"], default="common",
                   help="Extra microgrid meters (DB1, DB2, ...): 'common' = shared by UE like Common Services; "
                        "'ignore' = report only, charge nothing (reproduces the original spreadsheet)")
    o.add_argument("--new-asset-baseline", choices=["zero", "skip"], default="zero",
                   help="Extra/EV meter without a previous read: 'zero' = register starts at 0; 'skip' = leave out")
    o.add_argument("--swap-import-export", default="BESS",
                   help="Comma-separated assets whose CSV import/export are reversed (Stormcloud reports BESS "
                        "reversed); 'none' to disable")
    o.add_argument("--timezone", default="Australia/Perth", help="Timezone for assigning CSV periods to months")
    o.add_argument("--nmi", default=None, help="Supplier NMI to read (default: first 'Usage Calculations' block)")

    out = p.add_argument_group("Output")
    out.add_argument("--outdir", default="invoices", help="Output directory")
    out.add_argument("--issue-date", default=None, help="Invoice issue date YYYY-MM-DD (default: today)")
    out.add_argument("--due-days", type=int, default=14, help="Days until payment is due")
    out.add_argument("--title", default="Cluster Energy Billing", help="Invoice heading")
    out.add_argument("--no-auto-install", action="store_true",
                     help="Do not pip-install missing Python packages; just report them")
    out.add_argument("--issuer", default="The Owners of the Cluster",
                     help="Issuer line printed on invoices")
    x = p.add_argument_group("Xero export (invoices_xero_data.csv / credit_note_xero_data.csv)")
    x.add_argument("--no-xero", action="store_true", help="Do not write the Xero import CSVs")
    x.add_argument("--xero-contacts", default=os.path.join(here, "Lot_Names_and_Addresses.csv"),
                   help="Xero contacts export (AccountNumber LOTn -> *ContactName)")
    x.add_argument("--xero-invoice-suffix", default=None,
                   help="Suffix of the Xero invoice numbers, e.g. '/0326' gives 1A6/0326 (and 1A6/0326CN). "
                        "Must be unique per billing run; default '/MMYY' of the invoice date")
    x.add_argument("--xero-invoice-date", default=None, help="Xero *InvoiceDate YYYY-MM-DD (default: --issue-date)")
    x.add_argument("--xero-due-date", default=None, help="Xero *DueDate YYYY-MM-DD (default: invoice due date)")
    x.add_argument("--xero-account-charges", default="409", help="Account code: charges and fees")
    x.add_argument("--xero-account-imports", default="410", help="Account code: imports")
    x.add_argument("--xero-account-exports", default="411", help="Account code: exports")
    x.add_argument("--xero-tax-type", default="BAS Excluded", help="Xero *TaxType")
    x.add_argument("--xero-branding-theme", default="Energy Invoice", help="Xero BrandingTheme")
    return p


# -----------------------------------------------------------------------------
# Month helpers  (a month key is "YYYY-MM": the billing month a reading closes)
# -----------------------------------------------------------------------------
MONTH_ABB = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def parse_month_token(tok: str) -> int:
    t = tok.strip().lower()
    m = int(t) if t.isdigit() else (MONTH_ABB.index(t[:3]) + 1 if t[:3] in MONTH_ABB else 0)
    if not 1 <= m <= 12:
        raise BillingError(f"Cannot interpret '{tok}' as a month")
    return m


def parse_months(spec: str) -> list[int]:
    s = re.sub(r"\s*([-:])\s*", r"\1", spec)
    out: list[int] = []
    for part in filter(None, re.split(r"[,;\s]+", s)):
        if re.search(r"[-:]", part):
            ab = re.split(r"[-:]", part)
            a, b = parse_month_token(ab[0]), parse_month_token(ab[-1])
            if b < a:
                raise BillingError(f"Month range '{part}' runs backwards")
            out += list(range(a, b + 1))
        else:
            out.append(parse_month_token(part))
    if not out:
        raise BillingError("No billing months given (--months)")
    return sorted(set(out))


def add_months(key: str, n: int) -> str:
    y, m = int(key[:4]), int(key[5:7])
    t = m - 1 + n
    return f"{y + t // 12:04d}-{t % 12 + 1:02d}"


def key_date(key: str) -> dt.date:
    return dt.date(int(key[:4]), int(key[5:7]), 1)


def month_end(key: str) -> dt.date:
    return key_date(add_months(key, 1)) - dt.timedelta(days=1)


def mshort(key: str) -> str:
    return MONTH_ABB[int(key[5:7]) - 1].title()


def mlabel(key: str) -> str:
    return f"{mshort(key)} {key[:4]}"


def to_num(x) -> float:
    s = re.sub(r"[$,\s]", "", str(x))
    neg = "-" in s
    try:
        v = float(s.replace("-", ""))
    except ValueError:
        return float("nan")
    return -v if neg else v


def colkey(x: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(x).lower())


# -----------------------------------------------------------------------------
# 1. Households
# -----------------------------------------------------------------------------
def canon_lot(x: str) -> str:
    m = re.match(r"\s*Lot\s*0*([0-9]+)", str(x), re.I)
    return f"Lot {int(m.group(1)):02d}"


def load_households(path: str, battery_capacity: float, export_limit: float) -> pd.DataFrame:
    if not os.path.exists(path):
        raise BillingError(f"Households file not found: {path}")
    df = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    keys = {colkey(c): c for c in df.columns}

    def col(name, required=True, default=""):
        c = keys.get(colkey(name))
        if c is None:
            if required:
                raise BillingError(f"Households table is missing the column '{name}'")
            return pd.Series([default] * len(df), index=df.index)
        return df[c]

    lot = col("Lot").astype(str)
    df = df[lot.str.match(r"^\s*Lot\s*\d+", case=False, na=False)].copy()
    if df.empty:
        raise BillingError(f"No 'Lot NN' rows found in {path}")

    def txt(s):
        return s.fillna("").astype(str).str.strip().replace({"nan": ""})

    ue = pd.to_numeric(col("Unit Entitlement"), errors="coerce")
    if ue.isna().any():
        raise BillingError(f"Missing Unit Entitlement values in {path}")
    if abs(ue.sum() - 1000) > 1e-6:
        note(f"Unit entitlements sum to {ue.sum():g} (the workbook assumes 1000); UE % = UE / sum(UE).")
    pc = col("Post Code", False)
    pc = pc.apply(lambda v: f"{float(v):.0f}" if re.fullmatch(r"\s*\d+(\.0+)?\s*", str(v)) else ("" if pd.isna(v) else str(v)))

    h = pd.DataFrame({
        "Cluster": txt(col("Cluster")).values,
        "Lot": [canon_lot(v) for v in col("Lot")],
        "Block Type": txt(col("Block Type", False)).values,
        "Owners": txt(col("Owners")).values,
        "Street Address": txt(col("Street Address", False)).values,
        "Suburb": txt(col("Suburb", False)).values,
        "State": txt(col("State", False)).values,
        "Post Code": pc.values,
        "Unit Entitlement": ue.values,
        "Export Capacity (kW)": pd.to_numeric(col("Export Capacity (kW)", False, 0), errors="coerce").fillna(0).values,
        "Energy Meter Installed": txt(col("Energy Meter Installed")).str.upper().values,
        "PV Installed": txt(col("PV Installed", False)).str.upper().values,
        "Functional Droplet": txt(col("Functional Droplet", False)).str.upper().values,
    })
    if h["Lot"].duplicated().any():
        raise BillingError("Duplicate lots in the households table: " + ", ".join(h["Lot"][h["Lot"].duplicated()]))
    h["Identifier"] = h["Cluster"] + " " + h["Lot"]
    h["UE %"] = h["Unit Entitlement"] / h["Unit Entitlement"].sum()
    h["Battery Allocation"] = battery_capacity * h["UE %"]                       # Households!M
    h["Limited Export Capacity"] = h["Export Capacity (kW)"] * export_limit      # Households!P
    return h.sort_values("Lot").reset_index(drop=True)


# -----------------------------------------------------------------------------
# 2. Meter readings: one record = cumulative (import, export) at END of 'month'
# -----------------------------------------------------------------------------
CORE_ASSETS = ["Common Services", "BESS", "Grid", "EV Charger"]


def canon_asset(x: str) -> str:
    n = str(x).strip()
    n = re.sub(r"^[0-9]+[A-Za-z]?\s+", "", n)            # drop cluster code "1A "
    if re.match(r"^Lot\s*[0-9]+", n, re.I):
        return canon_lot(n)
    n = re.sub(r"^Strata\s*-\s*", "", n, flags=re.I)
    if re.search(r"BESS|Battery", n, re.I):
        return "BESS"
    if re.search(r"Common\s*Serv", n, re.I):
        return "Common Services"
    if re.search(r"^SMSB$|Grid", n, re.I):
        return "Grid"
    if re.search(r"\bEV\b", n, re.I):
        return "EV Charger"
    n = re.sub(r"^SMSB\s+", "", n, flags=re.I)
    n = re.sub(r"\s+Meter$", "", n, flags=re.I)
    return n.strip()


def read_meter_csv(path: str, swap: set[str], tz: str) -> pd.DataFrame:
    if not os.path.exists(path):
        raise BillingError(f"Meter CSV not found: {path}")
    df = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
    low = [str(c).lower() for c in df.columns]

    def pick(s):
        for i, c in enumerate(low):
            if s in c:
                return df.columns[i]
        raise BillingError(f"{os.path.basename(path)}: no column containing '{s}'")

    c_ps, c_imp, c_exp = pick("period start"), pick("import"), pick("export")
    name = df[df.columns[0]].fillna("").str.strip()
    df = df[name != ""]
    name = name[name != ""]
    ps = pd.to_datetime(df[c_ps], utc=True, errors="coerce")
    if ps.isna().any():
        raise BillingError(f"{os.path.basename(path)}: cannot parse 'Period Start' timestamps")
    month = ps.dt.tz_convert(tz).dt.strftime("%Y-%m")
    imp = df[c_imp].map(to_num)
    exp = df[c_exp].map(to_num)
    asset = name.map(canon_asset)
    sw = asset.isin(swap)
    imp2, exp2 = imp.where(~sw, exp), exp.where(~sw, imp)
    if month.nunique() > 1:
        note(f"{os.path.basename(path)} covers several months ({', '.join(month.unique())}); "
             "each row is assigned to the month of its own Period Start.")
    out = pd.DataFrame({"asset": asset.values, "label": name.values, "month": month.values,
                        "imp": imp2.values, "exp": exp2.values, "source": os.path.basename(path)})
    bad = out["imp"].isna() | out["exp"].isna()
    if bad.any():
        note(f"{os.path.basename(path)}: non-numeric reads ignored for {', '.join(out.label[bad])}")
    return out[~bad]


LEDGER_COLS = ["asset", "month", "imp", "exp", "label", "source"]


def read_ledger(path: str) -> pd.DataFrame:
    """A meter_readings_ledger.csv written by a previous run (already in cluster convention: no swap)."""
    df = pd.read_csv(path, dtype={"asset": str, "month": str, "label": str, "source": str}, encoding="utf-8-sig")
    df["month"] = df["month"].str.strip()
    if not df["month"].str.fullmatch(r"\d{4}-\d{2}").all():
        raise BillingError(f"{os.path.basename(path)}: 'month' must be YYYY-MM")
    df["imp"], df["exp"] = pd.to_numeric(df["imp"], errors="coerce"), pd.to_numeric(df["exp"], errors="coerce")
    for c in ("label", "source"):
        if c not in df:
            df[c] = os.path.basename(path)
    return df.dropna(subset=["imp", "exp"])[LEDGER_COLS]


def read_readings_any(path: str, swap: set[str], tz: str) -> pd.DataFrame:
    if not os.path.exists(path):
        raise BillingError(f"Readings file not found: {path}")
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise BillingError(f"{os.path.basename(path)}: workbooks are no longer read; pass meter_readings_ledger.csv "
                           "or the previous month's meter CSV instead")
    head = {c.strip().lower() for c in pd.read_csv(path, nrows=0, encoding="utf-8-sig").columns}
    if {"asset", "month", "imp", "exp"} <= head:
        return read_ledger(path)
    return read_meter_csv(path, swap, tz)


# -----------------------------------------------------------------------------
# 3. Supplier PDF ("Usage Calculations" -> items 1-11 of each month tab)
# -----------------------------------------------------------------------------
CAT_PATTERNS = {
    "on_peak": r"^On\s*-?\s*Peak\s*Energy",
    "off_peak": r"^Off\s*-?\s*Peak\s*Energy",
    "daily_supply": r"Daily\s*Supply",
    "ess": r"Essential\s*System",
    "solar_on": r"Solar\s*On\s*-?\s*Peak",
    "solar_off": r"Solar\s*Off\s*-?\s*Peak",
    "stp": r"\bSTP\b|Small[- ]?scale",
    "rtp": r"\bRTP\b|Renewable\s*(Energy\s*)?Target",
    "wp_network": r"Western\s*Power|Network\s*(Tariff|Charge)",
    "market_fee": r"Market\s*Fee",
}
CATEGORIES = list(CAT_PATTERNS) + ["other"]
FIXED_CATS = ["ess", "stp", "rtp", "wp_network", "market_fee", "other"]          # items 4,7,8,9,10,11
# Row layout of the workbook month tabs (item #, label, category)
TAB_ITEMS = [(1, "On peak energy", "on_peak"), (2, "Off peak energy", "off_peak"), (3, "Daily supply", "daily_supply"),
             (4, "Essential system services", "ess"), (5, "Solar on peak feedback", "solar_on"),
             (6, "Solar off-peak feedback", "solar_off"), (7, "Clean Energy STP Pass-Through", "stp"),
             (8, "Clean Energy RTP Pass-Through", "rtp"), (9, "Western Power AA5 Tariff Increase", "wp_network"),
             (10, "Market Fee Pass-Through", "market_fee"), (11, "Other", "other")]
CAT_LABEL = {"ess": "Essential System Services", "stp": "Clean Energy STP pass-through",
             "rtp": "Clean Energy RTP pass-through", "wp_network": "Network tariff adjustment",
             "market_fee": "AEMO market fee pass-through", "other": "Other supplier charges"}

AMOUNT_RE = re.compile(r"(-?\$?-?[0-9][0-9,]*\.[0-9]+)\s*$")
USAGE_RE = re.compile(r"([0-9][0-9,]*(?:\.[0-9]+)?)\s*(kWh|Days|Day|kVA|kW)\b")
RATE_RE = re.compile(r"([0-9][0-9,]*(?:\.[0-9]+)?)\s*\$\s*/\s*(kWh|Days|Day|kVA|kW)\b")


def pdf_text(path: str) -> str:
    if shutil.which("pdftotext"):
        txt = subprocess.run(["pdftotext", "-layout", "-enc", "UTF-8", path, "-"],
                             capture_output=True, text=True, check=True).stdout
    else:
        try:
            import pdfplumber
        except ImportError as e:
            raise BillingError("Reading PDFs needs 'pdfplumber' (pip install pdfplumber) or poppler 'pdftotext'") from e
        with pdfplumber.open(path) as pdf:
            txt = "\n".join((pg.extract_text(layout=True) or "") for pg in pdf.pages)
    for ch in "\u00ad\u2010\u2011\u2012\u2013\u2014\u2212":      # supplier prints minus as soft hyphen
        txt = txt.replace(ch, "-")
    return txt.replace("\u00a0", " ")


@dataclass
class Bill:
    file: str
    account: str
    month: str
    start: dt.date | None
    finish: dt.date | None
    days: float
    nmi: str
    lines: pd.DataFrame
    amount: dict
    usage: dict
    rate: dict
    subtotal: float
    gst: float
    total: float


def parse_charge_line(s: str):
    s = s.strip()
    am = AMOUNT_RE.search(s)
    if not am:
        return None
    body = s[:am.start()]
    us, rt = USAGE_RE.search(body), RATE_RE.search(body)
    cut = min([m.start() for m in (us, rt) if m] + [len(body)])
    label = body[:cut].strip()
    if not label:
        return None
    return {"label": label,
            "usage": to_num(us.group(1)) if us else np.nan, "usage_unit": us.group(2) if us else "",
            "rate": to_num(rt.group(1)) if rt else np.nan, "rate_unit": rt.group(2) if rt else "",
            "amount": to_num(am.group(1))}


def parse_supplier_pdf(path: str, nmi: str | None = None) -> Bill:
    if not os.path.exists(path):
        raise BillingError(f"Supplier PDF not found: {path}")
    base = os.path.basename(path)
    full = pdf_text(path)
    lines = full.splitlines()
    starts = [i for i, l in enumerate(lines) if "Usage Calculations" in l]
    if not starts:
        raise BillingError(f"{base}: no 'Usage Calculations' section found")
    nmis = []
    for s in starts:
        m = re.search(r"NMI\s*:?\s*([0-9][0-9 ]{8,16}[0-9])", " ".join(lines[s:s + 7]))
        nmis.append(re.sub(r"\D", "", m.group(1)) if m else None)
    pick = 0
    if nmi:
        want = re.sub(r"\D", "", nmi)
        if want not in nmis:
            raise BillingError(f"{base}: NMI {nmi} not found (blocks found: {nmis})")
        pick = nmis.index(want)
    s = starts[pick]
    ends = [i for i, l in enumerate(lines) if "Total This NMI" in l and i > s]
    if not ends:
        raise BillingError(f"{base}: 'Total This NMI' not found after 'Usage Calculations'")
    blk = lines[s:ends[0] + 1]
    hdr = [i for i, l in enumerate(blk) if re.search(r"^\s*(Charges|Description)\b.*Amount", l)]
    sub = [i for i, l in enumerate(blk) if re.search(r"Sub\s*Total", l)]
    gsti = [i for i, l in enumerate(blk) if "GST" in l]
    rows = []
    for l in blk[(hdr[0] + 1 if hdr else 1):(sub[0] if sub else len(blk) - 1)]:
        if not l.strip() or re.search(r"NMI\s*:|Supply Address", l):
            continue
        r = parse_charge_line(l)
        if r is None:
            note(f"{base}: could not read supplier line '{l.strip()}'")
            continue
        rows.append(r)
    if not rows:
        raise BillingError(f"{base}: no charge lines could be read in 'Usage Calculations'")
    ln = pd.DataFrame(rows)
    ln["category"] = [next((k for k, p in CAT_PATTERNS.items() if re.search(p, lb, re.I)), "other") for lb in ln.label]

    def amt_end(t):
        m = AMOUNT_RE.search(t.strip())
        return to_num(m.group(1)) if m else np.nan

    subtotal = amt_end(blk[sub[0]]) if sub else ln.amount.sum()
    gst = amt_end(blk[gsti[0]]) if gsti else np.nan
    total = amt_end(blk[-1])
    if np.isnan(total):
        raise BillingError(f"{base}: cannot read 'Total This NMI'")

    def grab(p):
        m = re.search(p, full)
        return m.group(1) if m else None

    def dmy(x):
        return dt.datetime.strptime(x, "%d/%m/%Y").date() if x else None

    start = dmy(grab(r"Start Time:\s*([0-9]{2}/[0-9]{2}/[0-9]{4})")) or \
        dmy(grab(r"Supply period:\s*([0-9]{2}/[0-9]{2}/[0-9]{4})"))
    finish = dmy(grab(r"Finish Time:\s*([0-9]{2}/[0-9]{2}/[0-9]{4})")) or \
        dmy(grab(r"Supply period:\s*[0-9]{2}/[0-9]{2}/[0-9]{4}\s*-\s*([0-9]{2}/[0-9]{2}/[0-9]{4})"))
    account = None
    m_acc = re.search(r"ACCOUNT REFERENCE NUMBER:", full, re.I)
    if m_acc:   # first token after the label containing letters AND >= 4 digits (skips addresses etc.)
        for tok in re.findall(r"[A-Za-z0-9]+", full[m_acc.end():m_acc.end() + 400]):
            if re.search(r"[A-Za-z]", tok) and len(re.findall(r"\d", tok)) >= 4:
                account = tok
                break
    account = account or os.path.splitext(base)[0]
    if start:
        month = start.strftime("%Y-%m")
    else:
        m = re.search(r"(20[0-9]{2})(0[1-9]|1[0-2])$", account)
        if not m:
            raise BillingError(f"{base}: cannot determine the supply period")
        month = f"{m.group(1)}-{m.group(2)}"

    amount = {k: 0.0 for k in CATEGORIES}
    usage = {k: 0.0 for k in CATEGORIES}
    rate = {k: np.nan for k in CATEGORIES}
    for r in ln.itertuples():
        amount[r.category] += r.amount
        if not np.isnan(r.usage):
            usage[r.category] += r.usage
        if not np.isnan(r.rate):
            rate[r.category] = r.rate
    if (ln.category == "other").any():
        note(f"{base}: unrecognised supplier line(s) {list(ln.label[ln.category == 'other'])} "
             "treated as supplier fixed fees (item 11).")
    if abs(ln.amount.sum() - subtotal) > 0.05:
        note(f"{base}: charge lines sum to {ln.amount.sum():.2f} but Sub Total is {subtotal:.2f} - check the parse.")
    if not np.isnan(gst) and abs(subtotal + gst - total) > 0.05:
        note(f"{base}: Sub Total + GST ({subtotal + gst:.2f}) differs from Total ({total:.2f}).")
    days = grab(r"Number of Days:\s*([0-9]+)")
    return Bill(base, account, month, start, finish, float(days) if days else np.nan, nmis[pick] or "",
                ln, amount, usage, rate, subtotal, gst, total)


# -----------------------------------------------------------------------------
# 4. One billing month (formulas of the "MM - Mon" tabs)
# -----------------------------------------------------------------------------
@dataclass
class MonthResult:
    month: str
    lots: pd.DataFrame
    strata: pd.DataFrame
    cluster: dict
    extras: list[str]
    comps: pd.DataFrame
    bill: Bill
    supplier: dict = field(default_factory=dict)


class Readings:
    def __init__(self, df: pd.DataFrame):
        self.df = df
        self.idx = {(a, m): (i, e) for a, m, i, e in zip(df.asset, df.month, df.imp, df.exp)}

    def get(self, asset, month):
        return self.idx.get((asset, month))


def calc_month(m: str, hh: pd.DataFrame, bill: Bill, rd: Readings, a) -> MonthResult:
    gx, gi, p2p, nh, g1 = a.grid_export_rate, a.grid_import_rate, a.p2p_rate, a.num_households, 1 + a.gst
    pm = add_months(m, -1)
    hint = (f"add the {mlabel(pm)} Stormcloud CSV to --meters (it provides the opening reads), "
            "or pass a previous run's meter_readings_ledger.csv via --opening-readings")

    def delta(asset, kind):
        """kind: required (lots, Common Services, BESS) | optional (EV, extra meters) | report (grid meter)"""
        cl, op = rd.get(asset, m), rd.get(asset, pm)
        if cl is None:
            if kind == "required":
                raise BillingError(f"No {mlabel(m)} closing read for '{asset}'.")
            return (np.nan, np.nan, False) if kind == "report" else (0.0, 0.0, False)
        if op is None:
            if kind == "required":
                raise BillingError(f"No opening read for '{asset}' in {mlabel(m)}: {hint}.")
            if kind == "report":
                return np.nan, np.nan, False
            if a.new_asset_baseline == "skip":
                note(f"{mlabel(m)}: '{asset}' has no {mlabel(pm)} read and is left out (--new-asset-baseline skip).")
                return 0.0, 0.0, False
            note(f"{mlabel(m)}: '{asset}' has no {mlabel(pm)} read; its register is assumed to start at 0 "
                 f"(cumulative {cl[0]:g}/{cl[1]:g} kWh billed in this month). Use --new-asset-baseline skip "
                 "or add the earlier CSV to change this.")
            op = (0.0, 0.0)
        d = (cl[0] - op[0], cl[1] - op[1])
        if min(d) < 0:
            note(f"{mlabel(m)}: '{asset}' register went backwards ({d[0]:g} / {d[1]:g} kWh) - meter reset?")
        return d[0], d[1], True

    A, U = bill.amount, bill.usage
    days = U["daily_supply"] if U["daily_supply"] > 0 else (bill.days if bill.days == bill.days and bill.days > 0
                                                           else (month_end(m) - key_date(m)).days + 1)
    energy_total = U["on_peak"] + U["off_peak"]                                              # D19
    solar_credit = U["solar_on"] + U["solar_off"]                                            # D20
    energy_charge = (A["on_peak"] + A["off_peak"] + A["solar_on"] + A["solar_off"]) * g1     # D21
    fixed_fees = sum(A[k] for k in FIXED_CATS) * g1                                          # D22
    daily_supply = A["daily_supply"] * g1                                                    # D23
    supplier_total = bill.total                                                              # E17

    cs = delta("Common Services", "required")
    bess = delta("BESS", "required")
    ev = delta("EV Charger", "optional")
    gm = delta("Grid", "report")
    cs_net, ev_net = cs[0] - cs[1], ev[0] - ev[1]                                             # K25, K27
    bess_eff = bess[1] / bess[0] if bess[0] > 0 else 0.0                                      # K29
    if bess[0] <= 0:
        note(f"{mlabel(m)}: BESS import is 0; BESS efficiency set to 0.")

    present = set(rd.df.asset[rd.df.month == m])
    extras = sorted(x for x in present if not re.fullmatch(r"Lot \d+", x) and x not in CORE_ASSETS)
    ex = {x: delta(x, "optional") for x in extras}
    ex_net = {x: v[0] - v[1] for x, v in ex.items()}
    stray = sorted(x for x in present if re.fullmatch(r"Lot \d+", x) and x not in set(hh.Lot))
    if stray:
        note(f"{mlabel(m)}: meter rows for {', '.join(stray)} are not in the households table and are ignored.")

    L = pd.DataFrame({
        "Cluster": hh["Cluster"], "Lot": hh["Lot"], "Owners": hh["Owners"],
        "Billing Month Year": key_date(m), "Identifier": hh["Identifier"],
        "Row ID": hh["Identifier"] + " " + key_date(m).strftime("%b-%y"),
        "Battery Allocation": hh["Battery Allocation"]})
    dl = [delta(l, "required") for l in hh.Lot]
    L["Total Export"] = [d[1] for d in dl]
    L["Total Import"] = [d[0] for d in dl]
    net_wp = solar_credit - energy_total                                                     # K30
    net_mg_export = L["Total Export"].sum() - solar_credit                                   # K32

    L["Total BA"] = L["Battery Allocation"] * days                                           # M
    L["Applied BA"] = np.minimum(L["Total Export"], L["Total BA"])                            # N
    s_aba = L["Applied BA"].sum()
    L["Applied BA %"] = L["Applied BA"] / s_aba if s_aba > 0 else 0.0                        # O
    L["MG Contribution"] = net_mg_export * L["Applied BA %"]                                  # P
    L["Grid Export"] = L["Total Export"] - L["MG Contribution"]                               # J
    L["MG Export"] = L["Total Export"] - L["Grid Export"]                                     # L
    s_gx = L["Grid Export"].sum()
    L["Grid Export %"] = L["Grid Export"] / s_gx if s_gx else 0.0                            # K
    L["Grid Export Credit"] = L["Grid Export"] * gx                                           # Q
    L["MG Export Credit"] = L["MG Export"] * p2p                                              # R
    s_ti = L["Total Import"].sum()
    L["Import %"] = L["Total Import"] / s_ti if s_ti > 0 else hh["UE %"]                     # T
    L["BESS Import Allocation"] = L["MG Contribution"] * bess_eff                             # U
    L["MG Import"] = np.minimum(L["BESS Import Allocation"], L["Total Import"])               # V
    L["Grid Import"] = L["Total Import"] - L["MG Import"]                                     # W
    s_gi = L["Grid Import"].sum()
    L["Grid Import %"] = L["Grid Import"] / s_gi if s_gi else 0.0                            # X
    L["Grid Import Variance"] = (energy_total - s_gi) * L["Grid Import %"]                    # Y (info)
    L["Grid Import Charge"] = L["Grid Import"] * gi                                           # Z
    L["MG Import Charge"] = L["MG Import"] * p2p                                              # AA
    L["Amanda Daily Supply"] = daily_supply / nh                                              # AB
    L["Energy Meter Installed"] = hh["Energy Meter Installed"]                                # AC
    L["Amanda Fixed Fees"] = np.where(hh["Energy Meter Installed"] == "Y", fixed_fees / nh, 0.0)  # AD
    common_rate = p2p if net_wp > 0 else gi
    L["Common Services Split"] = cs_net * hh["UE %"]                                          # AE
    L["Common Services Charge"] = L["Common Services Split"] * common_rate                    # AF
    L["EV Charger Split"] = ev_net * hh["UE %"]                                               # AG
    L["EV Charge"] = L["EV Charger Split"] * p2p                                              # AH
    charge_extras = a.extra_assets == "common"
    L["Extra Assets Charge"] = 0.0
    for x in extras:                               # NEW: DB1, DB2, ... shared like Common Services
        L[f"{x} Split"] = ex_net[x] * hh["UE %"]
        L[f"{x} Charge"] = L[f"{x} Split"] * common_rate if charge_extras else 0.0
        L["Extra Assets Charge"] += L[f"{x} Charge"]
    L["Invoice Before Var"] = (L["Grid Import Charge"] + L["MG Import Charge"] + L["Amanda Daily Supply"]
                               + L["Amanda Fixed Fees"] + L["Common Services Charge"] + L["EV Charge"]
                               + L["Extra Assets Charge"]) - (L["Grid Export Credit"] + L["MG Export Credit"])  # AJ
    var = supplier_total - L["Invoice Before Var"].sum()                                       # E25
    L["Variance"] = var * L["Import %"]                                                       # AI
    L["Final Invoice"] = L["Invoice Before Var"] + L["Variance"]                              # AK

    comps = []
    for k in FIXED_CATS:
        if A[k] != 0:
            comps.append(pd.DataFrame({"Lot": hh.Lot, "month": m, "category": k, "cluster_usage": U[k],
                                       "cluster_amount": A[k] * g1,
                                       "amount": np.where(hh["Energy Meter Installed"] == "Y", A[k] * g1 / nh, 0.0)}))
    comps = pd.concat(comps, ignore_index=True) if comps else pd.DataFrame(
        columns=["Lot", "month", "category", "cluster_usage", "cluster_amount", "amount"])

    strata_rows = [("Strata - Common Services", cs), ("Strata - BESS", bess), ("Strata - EV Charger", ev)] + \
                  [(f"Strata - {x}", ex[x]) for x in extras] + [("Strata - SMSB (grid meter)", gm)]
    strata = pd.DataFrame({"Cluster": hh.Cluster.iloc[0], "Lot": [n for n, _ in strata_rows],
                           "Billing Month Year": key_date(m),
                           "Total Export": [v[1] for _, v in strata_rows],
                           "Total Import": [v[0] for _, v in strata_rows]})

    metered_net = s_ti + cs_net + ev_net + sum(ex_net.values()) + bess[0] - bess[1] - L["Total Export"].sum()
    cl = dict(month=m, supplier_invoice=bill.account, supplier_total=supplier_total, supply_days=days,
              supplier_import_kwh=energy_total, supplier_solar_feedback_kwh=solar_credit, net_wp_imp_exp_kwh=net_wp,
              supplier_energy_charge=energy_charge, supplier_fixed_fees=fixed_fees, supplier_daily_supply=daily_supply,
              lots_import_kwh=s_ti, lots_export_kwh=L["Total Export"].sum(), net_mg_export_kwh=net_mg_export,
              mg_import_kwh=L["MG Import"].sum(), grid_import_kwh=s_gi, common_services_kwh=cs_net,
              ev_charger_kwh=ev_net, bess_import_kwh=bess[0], bess_export_kwh=bess[1], bess_efficiency=bess_eff,
              bess_loss_kwh=bess[0] - bess[1], grid_meter_import_kwh=gm[0], grid_meter_export_kwh=gm[1],
              metered_net_kwh=metered_net, metering_residual_kwh=(gm[0] - gm[1]) - metered_net,
              common_rate=common_rate, grid_import_variance_kwh=energy_total - s_gi, variance=var,
              invoice_before_var=L["Invoice Before Var"].sum(), final_total=L["Final Invoice"].sum())
    for x in extras:
        cl[f"extra_{x}_kwh"] = ex_net[x]
    return MonthResult(m, L, strata, cl, extras, comps, bill,
                       supplier=dict(energy_total=energy_total, solar_credit=solar_credit,
                                     energy_charge=energy_charge, fixed_fees=fixed_fees, daily_supply=daily_supply))


# -----------------------------------------------------------------------------
# 5a. Monthly tabs workbook
# -----------------------------------------------------------------------------
LOT_COLS = ["Cluster", "Lot", "Owners", "Billing Month Year", "Identifier", "Row ID", "Battery Allocation",
            "Total Export", "Grid Export", "Grid Export %", "MG Export", "Total BA", "Applied BA", "Applied BA %",
            "MG Contribution", "Grid Export Credit", "MG Export Credit", "Total Import", "Import %",
            "BESS Import Allocation", "MG Import", "Grid Import", "Grid Import %", "Grid Import Variance",
            "Grid Import Charge", "MG Import Charge", "Amanda Daily Supply", "Energy Meter Installed",
            "Amanda Fixed Fees", "Common Services Split", "Common Services Charge", "EV Charger Split", "EV Charge"]
TAIL_COLS = ["Variance", "Invoice Before Var", "Final Invoice"]
NO_TOTAL = {"Cluster", "Lot", "Owners", "Billing Month Year", "Identifier", "Row ID", "Energy Meter Installed",
            "Total BA", "Grid Import Variance"}


def write_monthly_tabs(path: str, results: list[MonthResult], hh: pd.DataFrame, a) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    hfill = PatternFill("solid", fgColor="1F3A5F")
    sfill = PatternFill("solid", fgColor="E9EFF6")
    newfill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="C9D3DE")
    bd = Border(bottom=thin)
    hdr_font = Font(name=F, bold=True, color="FFFFFF", size=9)
    base = Font(name=F, size=9)
    bold = Font(name=F, size=9, bold=True)
    title = Font(name=F, size=13, bold=True, color="1F3A5F")
    sec = Font(name=F, size=10, bold=True, color="1F3A5F")
    FMT_KWH, FMT_D, FMT_P, FMT_R = "#,##0.00;-#,##0.00", '"$"#,##0.00;-"$"#,##0.00', "0.00%", '"$"0.0000'

    wb = Workbook()
    wb.remove(wb.active)
    for r in results:
        ws = wb.create_sheet(f"{r.month[5:7]} - {mshort(r.month)}")
        ws.sheet_view.showGridLines = False
        ws["B2"] = key_date(r.month)
        ws["B2"].number_format = "mmmm yyyy"
        ws["B2"].font = title
        ws.merge_cells("B2:C2")
        ws["D2"] = f"Supplier invoice {r.bill.account} ({r.bill.file})"
        ws["D2"].font = Font(name=F, size=9, italic=True, color="5F6B76")

        # items 1-11 (supplier "Usage Calculations")
        for j, h in enumerate(["#", "Supplier item", "Usage", "Rate", "Amount ($)"]):
            c = ws.cell(row=3, column=2 + j, value=h)
            c.font, c.fill = hdr_font, hfill
        for i, (n, lbl, k) in enumerate(TAB_ITEMS):
            row = 4 + i
            ws.cell(row=row, column=2, value=n).font = base
            ws.cell(row=row, column=3, value=lbl).font = base
            u = r.bill.usage[k]
            if u:
                ws.cell(row=row, column=4, value=u).number_format = FMT_KWH
            if r.bill.rate[k] == r.bill.rate[k]:
                ws.cell(row=row, column=5, value=r.bill.rate[k]).number_format = "0.000000"
            if r.bill.amount[k]:
                ws.cell(row=row, column=6, value=r.bill.amount[k]).number_format = FMT_D
            for col in range(2, 7):
                ws.cell(row=row, column=col).font = base
                ws.cell(row=row, column=col).border = bd
        for i, (lbl, v) in enumerate([("Sub Total", r.bill.subtotal), ("Plus GST", r.bill.gst), ("Total", r.bill.total)]):
            ws.cell(row=15 + i, column=5, value=lbl).font = bold
            c = ws.cell(row=15 + i, column=6, value=v)
            c.number_format, c.font = FMT_D, bold

        sp = r.supplier
        for i, (lbl, v, fmt) in enumerate([
                ("Amanda Energy Total (1+2), kWh", sp["energy_total"], FMT_KWH),
                ("Amanda Solar Credit (5+6), kWh", sp["solar_credit"], FMT_KWH),
                ("Amanda Energy Charge (1+2+5+6 + GST)", sp["energy_charge"], FMT_D),
                ("Amanda Fixed Fees Charge (4+7+8+9+10+11 + GST)", sp["fixed_fees"], FMT_D),
                ("Amanda Daily Supply Charge (3 + GST)", sp["daily_supply"], FMT_D)]):
            ws.cell(row=19 + i, column=3, value=lbl).font = base
            c = ws.cell(row=19 + i, column=4, value=v)
            c.number_format, c.font = fmt, base
        ws.cell(row=25, column=3, value="Variance (supplier total - households before variance)").font = bold
        c = ws.cell(row=25, column=5, value=r.cluster["variance"])
        c.number_format, c.font = FMT_D, bold

        # cluster metrics (J25:K32 in the workbook) + extra meters
        cl = r.cluster
        metrics = [("Strata - Common Services (kWh)", cl["common_services_kwh"], FMT_KWH),
                   ("Strata - BESS loss (kWh)", cl["bess_loss_kwh"], FMT_KWH),
                   ("Strata - EV Charger (kWh)", cl["ev_charger_kwh"], FMT_KWH)]
        metrics += [(f"Strata - {x} (kWh){'' if a.extra_assets == 'common' else ' [not charged]'}",
                     cl[f"extra_{x}_kwh"], FMT_KWH) for x in r.extras]
        metrics += [("BESS Efficiency", cl["bess_efficiency"], FMT_P),
                    ("Net WP Imp-Exp (kWh)", cl["net_wp_imp_exp_kwh"], FMT_KWH),
                    ("Grid Import Variance (kWh)", cl["grid_import_variance_kwh"], FMT_KWH),
                    ("Net MG Export (kWh)", cl["net_mg_export_kwh"], FMT_KWH),
                    ("Common load rate ($/kWh)", cl["common_rate"], FMT_R),
                    ("Metering residual vs SMSB grid meter (kWh)", cl["metering_residual_kwh"], FMT_KWH)]
        ws.cell(row=3, column=10, value="Cluster metrics").font = sec
        for i, (lbl, v, fmt) in enumerate(metrics):
            ws.cell(row=4 + i, column=10, value=lbl).font = base
            c = ws.cell(row=4 + i, column=11, value=None if v != v else v)
            c.number_format, c.font = fmt, base
            if any(lbl.startswith(f"Strata - {x}") for x in r.extras):
                ws.cell(row=4 + i, column=10).fill = newfill
                c.fill = newfill

        # strata table
        top = max(28, 4 + len(metrics) + 2)
        ws.cell(row=top - 1, column=2, value="Strata / microgrid assets (kWh this month)").font = sec
        for j, h in enumerate(["Cluster", "Asset", "Billing Month Year", "Total Export", "Total Import"]):
            c = ws.cell(row=top, column=2 + j, value=h)
            c.font, c.fill = hdr_font, hfill
        for i, row in enumerate(r.strata.itertuples(index=False)):
            vals = [row[0], row[1], row[2], row[3], row[4]]
            for j, v in enumerate(vals):
                c = ws.cell(row=top + 1 + i, column=2 + j, value=None if isinstance(v, float) and v != v else v)
                c.font, c.border = base, bd
                if j == 2:
                    c.number_format = "mmm-yy"
                if j >= 3:
                    c.number_format = FMT_KWH
        # lots table
        extra_cols = [f"{x} {s}" for x in r.extras for s in ("Split", "Charge")]
        cols = LOT_COLS + extra_cols + TAIL_COLS
        t0 = top + len(r.strata) + 3
        ws.cell(row=t0 - 1, column=2, value="Households").font = sec
        for j, h in enumerate(cols):
            c = ws.cell(row=t0, column=2 + j, value=h)
            c.font, c.fill = hdr_font, (PatternFill("solid", fgColor="B7791F") if h in extra_cols else hfill)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.row_dimensions[t0].height = 30
        money_cols = {"Grid Export Credit", "MG Export Credit", "Grid Import Charge", "MG Import Charge",
                      "Amanda Daily Supply", "Amanda Fixed Fees", "Common Services Charge", "EV Charge",
                      "Variance", "Invoice Before Var", "Final Invoice"} | {f"{x} Charge" for x in r.extras}
        pct_cols = {"Grid Export %", "Applied BA %", "Import %", "Grid Import %"}
        for i, row in enumerate(r.lots[cols].itertuples(index=False)):
            for j, (h, v) in enumerate(zip(cols, row)):
                c = ws.cell(row=t0 + 1 + i, column=2 + j, value=v)
                c.font, c.border = base, bd
                c.number_format = ("mmm-yy" if h == "Billing Month Year" else FMT_D if h in money_cols
                                   else FMT_P if h in pct_cols else FMT_KWH if j >= 6 and h != "Energy Meter Installed"
                                   else "General")
        tr = t0 + 1 + len(r.lots)
        ws.cell(row=tr, column=2, value="TOTALS").font = bold
        for j, h in enumerate(cols):
            if h in NO_TOTAL or j < 6:
                continue
            c = ws.cell(row=tr, column=2 + j, value=float(r.lots[h].sum()))
            c.font, c.fill = bold, sfill
            c.number_format = FMT_D if h in money_cols else FMT_P if h in pct_cols else FMT_KWH
        widths = {2: 8, 3: 46, 4: 30, 5: 13, 6: 13, 7: 18}
        for col in range(2, 2 + len(cols)):
            ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 13)
        ws.column_dimensions["J"].width = 40          # cluster-metric labels live in column J
        ws.freeze_panes = ws.cell(row=t0 + 1, column=5)

    # stacked table ("Master - Households Data")
    ws = wb.create_sheet("Master - Households Data")
    allx = sorted({x for r in results for x in r.extras})
    cols = LOT_COLS + [f"{x} {s}" for x in allx for s in ("Split", "Charge")] + TAIL_COLS
    master = pd.concat([r.lots for r in results], ignore_index=True).reindex(columns=cols)
    for j, h in enumerate(cols):
        c = ws.cell(row=1, column=1 + j, value=h)
        c.font, c.fill = hdr_font, hfill
    for i, row in enumerate(master.itertuples(index=False)):
        for j, v in enumerate(row):
            c = ws.cell(row=2 + i, column=1 + j, value=None if isinstance(v, float) and v != v else v)
            c.font = base
            if cols[j] == "Billing Month Year":
                c.number_format = "mmm-yy"
    ws.freeze_panes = "G2"

    ws = wb.create_sheet("Notes")
    lines = ["Generated by cluster_energy_billing.py on " + dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
             "Values are computed by the script (not live Excel formulas); column letters follow the original tabs.",
             f"Extra microgrid meters: --extra-assets {a.extra_assets}, --new-asset-baseline {a.new_asset_baseline}. "
             "Their columns are highlighted.",
             f"Rates: P2P {a.p2p_rate}, grid import {a.grid_import_rate}, grid export {a.grid_export_rate}, "
             f"GST {a.gst}, households {a.num_households}, battery {a.battery_capacity} kWh."] + \
            [f"NOTE: {n}" for n in NOTES]
    for i, l in enumerate(lines):
        ws.cell(row=1 + i, column=1, value=l).font = base
    ws.column_dimensions["A"].width = 160
    wb.save(path)


# -----------------------------------------------------------------------------
# 5b. Invoices (reportlab, A4)
# -----------------------------------------------------------------------------
def ascii(s) -> str:
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()


def f2(x, d=2) -> str:
    return f"{x:,.{d}f}"


def money(x) -> str:
    x = round(float(x), 2) + 0.0
    return f"-${-x:,.2f}" if x < 0 else f"${x:,.2f}"


def price(x) -> str:
    if x is None or x != x or not np.isfinite(x):
        return ""
    ax = abs(x)
    s = f"{ax:.6f}" if 0 < ax < 0.01 else f"{ax:.4f}"
    return ("-$" if x < 0 else "$") + s


def pct(x, d=2) -> str:
    return f"{100 * x:.{d}f}%"


def build_items(lot: str, results: list[MonthResult], a, hrow) -> dict:
    nh = a.num_households
    rows = [r.lots[r.lots.Lot == lot].iloc[0] for r in results]

    def S(c):
        return float(sum(x.get(c, 0.0) for x in rows))

    def eff(amt, qty):
        return amt / qty if abs(qty) > 1e-9 else np.nan

    uep = pct(hrow["UE %"])
    items = []

    def add(desc, qty, unit, up, amount, kind, qty_txt=None):
        items.append(dict(desc=desc, qty=qty, unit=unit, price=up, amount=amount, kind=kind, qty_txt=qty_txt))

    add("Microgrid (P2P) energy import", S("MG Import"), "kWh", a.p2p_rate, S("MG Import Charge"), "charge")
    add("Main grid energy import", S("Grid Import"), "kWh", a.grid_import_rate, S("Grid Import Charge"), "charge")
    q, v = S("Common Services Split"), S("Common Services Charge")
    add(f"Common services share (UE {uep})", q, "kWh", eff(v, q), v, "charge")
    if a.extra_assets == "common":
        for x in sorted({x for r in results for x in r.extras}):
            q, v = S(f"{x} Split"), S(f"{x} Charge")
            add(f"{x} microgrid load share (UE {uep})", q, "kWh", eff(v, q), v, "charge")
    q, v = S("EV Charger Split"), S("EV Charge")
    if abs(q) > 1e-9 or abs(v) > 1e-9:
        add(f"EV charger share (UE {uep})", q, "kWh", eff(v, q), v, "charge")
    days = sum(r.cluster["supply_days"] for r in results)
    v = S("Amanda Daily Supply")
    add(f"Supplier daily supply charge (1/{nh} share)", days, "days", eff(v, days), v, "charge")
    comps = pd.concat([r.comps[r.comps.Lot == lot] for r in results], ignore_index=True)
    for k in [c for c in FIXED_CATS if c in set(comps.category)]:
        cc = comps[comps.category == k]
        amt, use = float(cc.amount.sum()), float(cc.cluster_usage.sum())
        lbl = f"Pass-through: {CAT_LABEL[k]} (1/{nh})"
        meter = hrow["Energy Meter Installed"] == "Y"
        if use > 0:
            qty = use / nh if meter else 0.0
            add(lbl, qty, "kWh", eff(amt, qty), amt, "charge")
        else:
            add(lbl, np.nan, "share", float(cc.cluster_amount.sum()), amt, "charge",
                qty_txt=f"1/{nh}" if meter else "0")
    add("Microgrid (P2P) export credit", S("MG Export"), "kWh", a.p2p_rate, -S("MG Export Credit"), "credit")
    add("Main grid export credit", S("Grid Export"), "kWh", a.grid_export_rate, -S("Grid Export Credit"), "credit")
    return dict(items=items, rows=rows, ibv=S("Invoice Before Var"), var=S("Variance"), ti=S("Total Import"),
                te=S("Total Export"), total=S("Final Invoice"))


def render_invoice(path: str, lot: str, results: list[MonthResult], a, hrow, ctx) -> float:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    NAVY, GREEN, LIGHT, GREY = (colors.HexColor(c) for c in ("#1F3A5F", "#3C8D5A", "#E9EFF6", "#5F6B76"))
    RED = colors.HexColor("#C0392B")
    LINE, SHADE = colors.HexColor("#C9D3DE"), colors.HexColor("#F5F7FA")
    b = build_items(lot, results, a, hrow)
    if abs(sum(i["amount"] for i in b["items"]) - b["ibv"]) > 1e-6:
        note(f"{lot}: itemised lines differ from the calculated subtotal")
    mk = [r.month for r in results]
    inv_no = f"{hrow['Cluster']}-{lot.replace(' ', '')}-{ctx['period_code']}"
    W, H = A4
    CW = W - 24 * mm

    st_n = ParagraphStyle("n", fontName="Helvetica", fontSize=7.3, leading=9)
    st_s = ParagraphStyle("s", parent=st_n, fontSize=7, leading=9, textColor=GREY)
    st_h = ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=NAVY, spaceBefore=5,
                          spaceAfter=2)
    st_hc = ParagraphStyle("hc", fontName="Helvetica-Bold", fontSize=6.8, leading=8, textColor=NAVY, alignment=2)
    st_hl = ParagraphStyle("hl", parent=st_hc, alignment=0)

    def on_page(c, doc):
        c.saveState()
        c.setFillColor(NAVY)
        c.rect(0, H - 22 * mm, W, 22 * mm, stroke=0, fill=1)
        c.setFillColor(GREEN)
        c.rect(0, H - 23.2 * mm, W, 1.2 * mm, stroke=0, fill=1)
        c.setFillColor(colors.white)
        c.setFont("Helvetica-Bold", 13)
        c.drawString(12 * mm, H - 10 * mm, ascii(ctx["title"]))
        c.drawRightString(W - 12 * mm, H - 10 * mm, "ENERGY INVOICE")
        c.setFont("Helvetica", 7.5)
        c.drawString(12 * mm, H - 17 * mm, ascii(ctx["issuer"]))
        c.drawRightString(W - 12 * mm, H - 17 * mm, f"{inv_no}  |  {hrow['Identifier']}")
        c.setStrokeColor(GREEN)
        c.setLineWidth(1.2)
        c.line(12 * mm, 13 * mm, W - 12 * mm, 13 * mm)
        c.setFillColor(GREY)
        c.setFont("Helvetica", 6.5)
        c.drawString(12 * mm, 9 * mm, ascii(f"Generated {dt.date.today():%d/%m/%Y} by cluster_energy_billing.py  |  "
                                            f"Supplier invoice(s) {ctx['supplier_short']}  |  Stormcloud meter data"))
        c.drawRightString(W - 12 * mm, 9 * mm, f"Page {doc.page}")
        c.restoreState()

    def table(header, data, widths, bold_rows=(), shade_rows=(), italic_rows=(), align_left=(0,), size=7.3):
        hdr = [Paragraph(h.replace("\n", "<br/>"), st_hl if j in align_left else st_hc) for j, h in enumerate(header)]
        t = Table([hdr] + data, colWidths=[w * mm for w in widths], repeatRows=1, hAlign="LEFT")
        sty = [("FONT", (0, 1), (-1, -1), "Helvetica", size), ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
               ("ALIGN", (0, 0), (-1, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
               ("LINEBELOW", (0, 1), (-1, -1), 0.3, LINE), ("TOPPADDING", (0, 0), (-1, -1), 1.6),
               ("BOTTOMPADDING", (0, 0), (-1, -1), 1.6), ("LEFTPADDING", (0, 0), (-1, -1), 3),
               ("RIGHTPADDING", (0, 0), (-1, -1), 3)]
        for j in align_left:
            sty.append(("ALIGN", (j, 0), (j, -1), "LEFT"))
        for r in bold_rows:
            sty += [("FONT", (0, r + 1), (-1, r + 1), "Helvetica-Bold", size),
                    ("LINEABOVE", (0, r + 1), (-1, r + 1), 0.7, NAVY)]
        for r in shade_rows:
            sty.append(("BACKGROUND", (0, r + 1), (-1, r + 1), SHADE))
        for r in italic_rows:
            sty += [("FONT", (0, r + 1), (-1, r + 1), "Helvetica-Oblique", size),
                    ("TEXTCOLOR", (0, r + 1), (-1, r + 1), GREY)]
        t.setStyle(TableStyle(sty))
        return t

    story = []
    # ---- parties & details ---------------------------------------------------
    left = [Paragraph("<font size=7 color='#5F6B76'><b>BILL TO</b></font>", st_n),
            Paragraph(f"<font size=10><b>{ascii(hrow['Owners'])}</b></font>", ParagraphStyle("o", leading=13)),
            Paragraph(ascii(hrow["Street Address"]), ParagraphStyle("a", fontSize=8.5, leading=11)),
            Paragraph(ascii(f"{hrow['Suburb']} {hrow['State']} {hrow['Post Code']}".strip()),
                      ParagraphStyle("a2", fontSize=8.5, leading=11)),
            Paragraph(ascii(hrow["Identifier"] + (f"  ({hrow['Block Type']})" if hrow["Block Type"] else "")),
                      ParagraphStyle("a3", fontSize=8.5, leading=11, textColor=GREY))]
    det = [["Invoice no.", inv_no], ["Issue date", f"{ctx['issue']:%d/%m/%Y}"], ["Due date", f"{ctx['due']:%d/%m/%Y}"],
           ["Billing period", f"{key_date(mk[0]):%d/%m/%Y} - {month_end(mk[-1]):%d/%m/%Y}"],
           ["Months", ctx["months_txt"]], ["Supplier inv.", ctx["supplier_short"]]]
    dt_tab = Table([[Paragraph("<font size=7 color='#5F6B76'><b>INVOICE DETAILS</b></font>", st_n), ""]] + det,
                   colWidths=[24 * mm, 62 * mm])
    dt_tab.setStyle(TableStyle([("FONT", (0, 1), (0, -1), "Helvetica", 7.8), ("TEXTCOLOR", (0, 1), (0, -1), GREY),
                                ("FONT", (1, 1), (1, -1), "Helvetica-Bold", 7.8), ("SPAN", (0, 0), (1, 0)),
                                ("TOPPADDING", (0, 0), (-1, -1), 0.8), ("BOTTOMPADDING", (0, 0), (-1, -1), 0.8),
                                ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    top = Table([[left, dt_tab]], colWidths=[CW - 88 * mm, 88 * mm])
    top.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
    story += [top, Spacer(1, 4 * mm)]

    # ---- amount box ------------------------------------------------------------
    tot = b["total"]
    credit = round(tot, 2) <= 0
    accent = GREEN if credit else NAVY
    box = Table([[Paragraph(f"<b>{'CREDIT FOR THIS PERIOD' if credit else 'AMOUNT DUE FOR THIS PERIOD'}</b><br/>"
                            f"<font size=7.5 color='#5F6B76'>"
                            f"{'This credit will be applied to your account.' if credit else 'Please pay by ' + format(ctx['due'], '%d/%m/%Y') + '.'}"
                            f"</font>", ParagraphStyle("b", fontSize=8, leading=12, textColor=NAVY)),
                  Paragraph(f"<b>{money(abs(tot))}</b>", ParagraphStyle("bb", fontSize=18, leading=22, alignment=2,
                                                                         textColor=GREEN if credit else RED))]],
                colWidths=[CW - 60 * mm, 60 * mm])
    box.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), LIGHT), ("LINEBEFORE", (0, 0), (0, 0), 4, accent),
                             ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 6),
                             ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LEFTPADDING", (0, 0), (0, 0), 8),
                             ("RIGHTPADDING", (1, 0), (1, 0), 8)]))
    story += [box, Spacer(1, 2 * mm),
              Paragraph(f"Unit entitlement {hrow['Unit Entitlement']:g} ({pct(hrow['UE %'])})  |  Battery allocation "
                        f"{hrow['Battery Allocation']:.3f} kWh/day  |  Export capacity {hrow['Export Capacity (kW)']:.1f} kW "
                        f"(limited to {hrow['Limited Export Capacity']:.2f} kW)  |  PV installed: "
                        f"{hrow['PV Installed'] or '-'}  |  Energy meter: {hrow['Energy Meter Installed'] or '-'}", st_s)]

    rows = b["rows"]

    def tot_row(lbl, cols, fmt):
        return [lbl] + [fmt(sum(x[c] for x in rows)) for c in cols]

    # ---- Invoice Gen: energy ---------------------------------------------------
    ecols = ["MG Export", "Grid Export", "Total Export", "MG Import", "Grid Import", "Total Import"]
    data = [[mlabel(m)] + [f2(x[c]) for c in ecols] for m, x in zip(mk, rows)] + [tot_row("Total", ecols, f2)]
    story += [Paragraph("Energy exported and imported (kWh)", st_h),
              table(["Month", "Microgrid\nexport", "Main grid\nexport", "TOTAL\nEXPORTED", "Microgrid\nimport",
                     "Main grid\nimport", "TOTAL\nIMPORTED"], data, [30] + [26] * 6, bold_rows=[len(data) - 1])]

    # ---- itemised (expanded Invoice Gen) -------------------------------------
    it = b["items"]

    def fq(i):
        if i["qty_txt"] is not None:
            return i["qty_txt"]
        if i["qty"] != i["qty"]:
            return ""
        return f2(i["qty"], 0) if i["unit"] == "days" else f2(i["qty"])

    def fp(i):
        if i["unit"] == "share":
            return money(i["price"]) + " total"
        p = price(i["price"])
        return f"{p}/{'day' if i['unit'] == 'days' else i['unit']}" if p else ""

    data, bold_r, shade_r, ital_r = [], [], [], []

    def push(v, bo=False, sh=False, itl=False):
        data.append(v)
        n = len(data) - 1
        if bo:
            bold_r.append(n)
        if sh:
            shade_r.append(n)
        if itl:
            ital_r.append(n)

    ch = [i for i in it if i["kind"] == "charge"]
    cr = [i for i in it if i["kind"] == "credit"]
    push(["Charges", "", "", "", ""], sh=True, itl=True)
    for i in ch:
        push([Paragraph(ascii(i["desc"]), st_n), fq(i), i["unit"], fp(i), money(i["amount"])])
    push(["Total charges", "", "", "", money(sum(i["amount"] for i in ch))], bo=True)
    push(["Credits", "", "", "", ""], sh=True, itl=True)
    for i in cr:
        push([Paragraph(ascii(i["desc"]), st_n), fq(i), i["unit"], fp(i), money(i["amount"])])
    push(["Total credits", "", "", "", money(sum(i["amount"] for i in cr))], bo=True)
    push(["Subtotal before supplier reconciliation", "", "", "", money(b["ibv"])], bo=True)
    push([Paragraph("Supplier bill reconciliation (allocated by import share)", st_n), f2(b["ti"]), "kWh",
          (price(b["var"] / b["ti"]) + "/kWh") if b["ti"] > 0 else "", money(b["var"])])
    push(["TOTAL FOR PERIOD", "", "", "", money(tot)], bo=True)
    story += [Paragraph(f"Itemised charges and credits <font size=7 color='#5F6B76'><i>&nbsp;&nbsp;summed over "
                        f"{len(mk)} month(s); unit prices are period averages</i></font>", st_h),
              table(["Description", "Quantity", "Unit", "Unit price", "Amount"], data, [88, 24, 14, 30, 30],
                    bold_rows=bold_r, shade_rows=shade_r, italic_rows=ital_r, align_left=(0, 2))]
    if b["ti"] > 0:
        story.append(Paragraph(f"Effective net cost: {price(tot / b['ti'])} per kWh imported (total {money(tot)} over "
                               f"{f2(b['ti'])} kWh imported, after {f2(b['te'])} kWh exported).", st_s))
    story.append(Paragraph(f"Rates applied: microgrid P2P {price(a.p2p_rate)}/kWh; main grid import "
                           f"{price(a.grid_import_rate)}/kWh; main grid export credit {price(a.grid_export_rate)}/kWh. "
                           f"Supplier-derived charges include GST at {pct(a.gst, 0)}.", st_s))

    # ---- Invoice Gen: credits ---------------------------------------------------
    ccols = ["MG Export Credit", "Grid Export Credit"]
    data = [[mlabel(m)] + [money(x[c]) for c in ccols] + [money(x["MG Export Credit"] + x["Grid Export Credit"])]
            for m, x in zip(mk, rows)]
    data.append(tot_row("Total", ccols, money) + [money(sum(x["MG Export Credit"] + x["Grid Export Credit"] for x in rows))])
    story.append(KeepTogether([Paragraph("Credits by month ($)", st_h),
                               table(["Month", "Microgrid export credit", "Main grid export credit", "Total credits"],
                                     data, [42, 48, 48, 48], bold_rows=[len(data) - 1])]))

    # ---- Invoice Gen: charges ---------------------------------------------------
    def common(x):
        return x["Common Services Charge"] + x.get("Extra Assets Charge", 0.0)

    data = [[mshort(m), f2(x["MG Import Charge"]), f2(x["Grid Import Charge"]), f2(x["Amanda Daily Supply"]),
             f2(x["Amanda Fixed Fees"]), f2(common(x)), f2(x["EV Charge"]), f2(-x["MG Export Credit"] - x["Grid Export Credit"]),
             f2(x["Variance"]), f2(x["Final Invoice"])] for m, x in zip(mk, rows)]
    data.append(["Total"] + [f2(sum(f(x) for x in rows)) for f in (
        lambda x: x["MG Import Charge"], lambda x: x["Grid Import Charge"], lambda x: x["Amanda Daily Supply"],
        lambda x: x["Amanda Fixed Fees"], common, lambda x: x["EV Charge"],
        lambda x: -x["MG Export Credit"] - x["Grid Export Credit"], lambda x: x["Variance"],
        lambda x: x["Final Invoice"])])
    story.append(KeepTogether([Paragraph("Charges by month ($)", st_h),
                               table(["Month", "Microgrid\nimport", "Main grid\nimport", "Daily\nsupply",
                                      "Fixed\nfees", "Common\nloads*", "EV energy\nsupply", "Export\ncredits",
                                      "Variance\napplied", "TOTAL CREDIT/\nCHARGE"], data,
                                     [16, 17, 17, 17, 17, 18, 18, 18, 18, 30], bold_rows=[len(data) - 1], size=6.9),
                               Paragraph("* Common services" + (" plus " + ", ".join(sorted({x for r in results for x in r.extras}))
                                         if a.extra_assets == "common" and any(r.extras for r in results) else "") +
                                         ", shared by unit entitlement.", st_s)]))

    # ---- how the shares were calculated ----------------------------------------
    data = [[mshort(r.month), f"{r.cluster['supply_days']:.0f}", f2(x["Total BA"], 1), f2(x["Applied BA"], 1),
             pct(x["Applied BA %"]), f2(x["MG Contribution"], 1), pct(r.cluster["bess_efficiency"], 1),
             pct(x["Import %"]), pct(hrow["UE %"]), price(r.cluster["common_rate"])] for r, x in zip(results, rows)]
    story.append(KeepTogether([Paragraph("How your share was calculated", st_h),
                               table(["Month", "Supply\ndays", "Battery\nalloc. kWh", "Applied\nalloc. kWh",
                                      "Applied\nalloc. share", "Microgrid\ncontrib. kWh", "BESS\nefficiency",
                                      "Import\nshare", "UE\nshare", "Common load\nrate $/kWh"], data,
                                     [16, 15, 19, 19, 19, 20, 18, 19, 17, 24], size=6.9)]))

    # ---- cluster context (chunks of 6 months) ---------------------------------
    xn = sorted({x for r in results for x in r.extras})
    for k0 in range(0, len(results), 6):
        rc = results[k0:k0 + 6]

        def met(lbl, f):
            return [lbl] + [f(r.cluster) for r in rc]

        nan = lambda v: v != v  # noqa: E731
        data = [met("Supplier invoice", lambda c: c["supplier_invoice"]),
                met("Supplier bill total (inc GST)", lambda c: money(c["supplier_total"])),
                met("Supplier grid import (kWh)", lambda c: f2(c["supplier_import_kwh"])),
                met("Supplier solar feedback (kWh)", lambda c: f2(c["supplier_solar_feedback_kwh"])),
                met("Households: total import (kWh)", lambda c: f2(c["lots_import_kwh"])),
                met("Households: total export (kWh)", lambda c: f2(c["lots_export_kwh"])),
                met("Net microgrid export pool (kWh)", lambda c: f2(c["net_mg_export_kwh"])),
                met("BESS charge / discharge (kWh)", lambda c: f"{c['bess_import_kwh']:,.0f} / {c['bess_export_kwh']:,.0f}"),
                met("Common services (kWh)", lambda c: f2(c["common_services_kwh"]))]
        for x in xn:
            data.append(met(f"{x} (kWh){'' if a.extra_assets == 'common' else ' - not charged'}",
                            lambda c, x=x: f2(c[f"extra_{x}_kwh"]) if f"extra_{x}_kwh" in c else "-"))
        data += [met("EV charger (kWh)", lambda c: f2(c["ev_charger_kwh"])),
                 met("SMSB grid meter import / export (kWh)", lambda c: "-" if nan(c["grid_meter_import_kwh"]) else
                     f"{c['grid_meter_import_kwh']:,.0f} / {c['grid_meter_export_kwh']:,.0f}"),
                 met("Metering residual (kWh)", lambda c: "-" if nan(c["metering_residual_kwh"]) else
                     f2(c["metering_residual_kwh"])),
                 met("Reconciliation allocated ($)", lambda c: money(c["variance"]))]
        cw = min(26, (186 - 66) / len(rc))
        story.append(KeepTogether([Paragraph("Cluster context (whole cluster, per month)", st_h),
                                   table([""] + [mlabel(r.month) for r in rc], data, [66] + [cw] * len(rc), size=6.9)]))

    # ---- method ----------------------------------------------------------------
    extra_txt = (", " + ", ".join(xn)) if (xn and a.extra_assets == "common") else ""
    story += [Paragraph("Method", st_h)] + [Paragraph(t, ParagraphStyle("m", parent=st_n, fontSize=7, spaceAfter=2)) for t in [
        "<b>Meter data:</b> monthly kWh are the differences between consecutive cumulative register reads from the "
        "Stormcloud meter exports.",
        "<b>Exports:</b> the cluster's net microgrid export pool (household exports minus the supplier's solar feedback) "
        "is shared in proportion to each household's applied battery allocation (the lesser of its export and its UE "
        f"share of the {a.battery_capacity:g} kWh battery x supply days). That share is credited at the P2P rate; "
        "the rest of your export is credited at the main grid export rate.",
        "<b>Imports:</b> your microgrid import is your microgrid contribution x BESS round-trip efficiency, capped at "
        "your total import, and is charged at the P2P rate. The remainder is main grid import at the grid import rate.",
        "<b>Supplier fixed costs:</b> the daily supply charge and pass-through fees on the supplier invoice (incl. GST) "
        f"are shared equally across {a.num_households} households; pass-through fees apply only where an energy meter "
        f"is installed. Common microgrid loads (common services{extra_txt}) are shared by unit entitlement, at the P2P "
        "rate when the cluster is a net grid exporter and at the grid import rate otherwise.",
        "<b>Reconciliation:</b> any difference between the supplier invoice and the sum of all household charges is "
        "allocated by each household's share of total household import, so the household invoices add up exactly to "
        "the supplier invoice."]]

    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=12 * mm, rightMargin=12 * mm, topMargin=28 * mm,
                            bottomMargin=18 * mm, title=f"{inv_no} {ascii(hrow['Owners'])}",
                            author=ascii(ctx["issuer"]))
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return tot


# -----------------------------------------------------------------------------
# 5c. Xero import CSVs (port of xero_v2_2026.R)
# -----------------------------------------------------------------------------
XERO_COLS = ["*ContactName", "EmailAddress", "POAddressLine1", "POAddressLine2", "POAddressLine3", "POAddressLine4",
             "POCity", "PORegion", "POPostalCode", "POCountry", "*InvoiceNumber", "Reference", "*InvoiceDate",
             "*DueDate", "InventoryItemCode", "*Description", "*Quantity", "*UnitAmount", "Discount", "*AccountCode",
             "*TaxType", "TrackingName1", "TrackingOption1", "TrackingName2", "TrackingOption2", "Currency",
             "BrandingTheme"]
XERO_IMPORT_COLS = ["MG Import Charge", "Grid Import Charge"]
XERO_EXPORT_COLS = ["MG Export Credit", "Grid Export Credit"]
XERO_FEE_COLS = ["Amanda Daily Supply", "Amanda Fixed Fees", "Common Services Charge", "EV Charge",
                 "Extra Assets Charge", "Variance"]          # EV + extra microgrid meters: new vs the R script


def load_xero_contacts(path: str, hh: pd.DataFrame) -> dict[str, str]:
    names = {}
    if os.path.exists(path):
        c = pd.read_csv(path, dtype=str, encoding="utf-8-sig")
        if not {"*ContactName", "AccountNumber"} <= set(c.columns):
            raise BillingError(f"{os.path.basename(path)} needs '*ContactName' and 'AccountNumber' columns")
        for acc, nm in zip(c["AccountNumber"], c["*ContactName"]):
            m = re.search(r"(\d+)", str(acc))
            if m and isinstance(nm, str):
                names[f"Lot {int(m.group(1)):02d}"] = nm.strip()
    else:
        note(f"Xero contacts file {path} not found; using household owner names as Xero contact names.")
    for lot, own in zip(hh.Lot, hh.Owners):
        if lot not in names:
            if os.path.exists(path):
                note(f"{lot} is not in {os.path.basename(path)}; using '{own}' as its Xero contact name.")
            names[lot] = own
    return names


def _csv_cell(v) -> str:
    if v is None or (isinstance(v, float) and v != v) or v == "":
        return ""
    if isinstance(v, (int, np.integer)):
        return str(int(v))
    return '"' + str(v).replace('"', '""') + '"'


def _write_xero(path: str, rows: list[dict]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        fh.write(",".join(f'"{c}"' for c in XERO_COLS) + "\n")
        for r in rows:
            fh.write(",".join(_csv_cell(r.get(c)) for c in XERO_COLS) + "\n")


def write_xero_csvs(od: str, results: list[MonthResult], hh: pd.DataFrame, a, ctx) -> tuple[str, str, int, int]:
    inv_date = dt.date.fromisoformat(a.xero_invoice_date) if a.xero_invoice_date else ctx["issue"]
    due_date = dt.date.fromisoformat(a.xero_due_date) if a.xero_due_date else ctx["due"]
    suffix = a.xero_invoice_suffix if a.xero_invoice_suffix is not None else f"/{inv_date:%m%y}"
    if a.xero_invoice_suffix is None:
        note(f"Xero invoice numbers use the suffix '{suffix}' (from the invoice date); pass --xero-invoice-suffix "
             "to choose another. Numbers must not repeat an earlier billing run.")
    contacts = load_xero_contacts(a.xero_contacts, hh)
    first, last = key_date(results[0].month), month_end(results[-1].month)
    period = f"for period {first.day} {first:%b %Y} to {last.day} {last:%b %Y}"
    detail = pd.concat([r.lots for r in results], ignore_index=True)
    for c in XERO_FEE_COLS:
        if c not in detail:
            detail[c] = 0.0
    sums = detail.groupby("Lot")[XERO_IMPORT_COLS + XERO_EXPORT_COLS + XERO_FEE_COLS + ["Final Invoice"]].sum()

    invoices, credits = [], []
    for h in hh.to_dict("records"):
        lot = h["Lot"]
        r = sums.loc[lot]
        fees, imp, exp, tot = (float(r[XERO_FEE_COLS].sum()), float(r[XERO_IMPORT_COLS].sum()),
                               float(r[XERO_EXPORT_COLS].sum()), float(r["Final Invoice"]))
        if abs(fees + imp - exp - tot) > 1e-6:
            raise BillingError(f"{lot}: Export/Import/Fee sum does not add up with the invoice total.")
        # round each line, letting the fees line absorb the rounding so the Xero total equals the PDF total
        tot_r, imp_r, exp_r = round(tot, 2), round(imp, 2), round(exp, 2)
        fees_r = round(tot_r - imp_r + exp_r, 2)
        n = int(re.search(r"\d+", lot).group())
        inv_no = f"{h['Cluster']}{n}{suffix}"
        base = {"*ContactName": contacts[lot], "*InvoiceNumber": inv_no, "Reference": f"{n:02d}LOT",
                "*InvoiceDate": f"{inv_date:%d/%m/%Y}", "*DueDate": f"{due_date:%d/%m/%Y}", "*Quantity": 1,
                "*TaxType": a.xero_tax_type, "BrandingTheme": a.xero_branding_theme}

        def line(desc, amount, acct):
            acct = int(acct) if str(acct).isdigit() else acct
            return {**base, "*Description": f"{desc} {period}", "*UnitAmount": f"{amount + 0.0:.2f}", "*AccountCode": acct}

        l_fee = line("Charges and Fees", fees_r, a.xero_account_charges)
        l_imp = line("Total Imports", imp_r, a.xero_account_imports)
        l_exp = line("Total Exports", -exp_r, a.xero_account_exports)
        if tot < 0:                      # credit: charges+imports invoiced, exports raised as a credit note
            invoices += [l_fee, l_imp]
            credits.append({**l_exp, "*InvoiceNumber": inv_no + "CN"})
        else:
            invoices += [l_fee, l_imp, l_exp]
    # same ordering as the R script: charge-lot invoices first, then credit-lot invoices
    charge_lots = [r for r in invoices if not any(c["Reference"] == r["Reference"] for c in credits)]
    credit_lots = [r for r in invoices if any(c["Reference"] == r["Reference"] for c in credits)]
    p_inv = os.path.join(od, "invoices_xero_data.csv")
    p_cn = os.path.join(od, "credit_note_xero_data.csv")
    _write_xero(p_inv, charge_lots + credit_lots)
    _write_xero(p_cn, credits)
    return p_inv, p_cn, len(charge_lots + credit_lots), len(credits)


# -----------------------------------------------------------------------------
# Main
# -----------------------------------------------------------------------------
def run(argv=None) -> list[MonthResult]:
    a = build_parser().parse_args(argv)
    if not 1 <= len(a.meters) <= 13:
        raise BillingError(f"Between 1 and 13 meter CSVs are expected (got {len(a.meters)}): one per target month "
                           "plus the month before the first target month")
    if a.num_households <= 0:
        raise BillingError("--num-households must be positive")
    if a.no_pv_p2p_rate is not None:
        note("--no-pv-p2p-rate is accepted for parity with Global Params but, as in the workbook, no formula uses it.")
    mkeys = [f"{a.year:04d}-{m:02d}" for m in parse_months(a.months)]
    issue = dt.date.fromisoformat(a.issue_date) if a.issue_date else dt.date.today()
    swap = set() if a.swap_import_export.lower() == "none" else {canon_asset(x) for x in a.swap_import_export.split(",")}

    info("Billing months: " + ", ".join(mlabel(m) for m in mkeys))
    hh = load_households(a.households, a.battery_capacity, a.export_limit)
    info(f"Households: {len(hh)} lots from {os.path.basename(a.households)}")
    if len(hh) != a.num_households:
        note(f"Households table has {len(hh)} lots but --num-households is {a.num_households}; supplier fixed costs "
             f"are divided by {a.num_households} (as in the workbook).")

    frames = [read_readings_any(f, swap, a.timezone) for f in a.opening_readings]
    for f in a.meters:
        x = read_meter_csv(f, swap, a.timezone)
        info(f"Meter CSV {os.path.basename(f)}: {len(x)} meters, month(s) {', '.join(x.month.unique())}")
        frames.append(x)
    rd_df = pd.concat(frames, ignore_index=True)
    rd_df = rd_df[~rd_df.duplicated(["asset", "month"], keep="last")].reset_index(drop=True)   # later inputs win
    rd = Readings(rd_df)

    bills = {}
    for f in a.pdfs:
        b = parse_supplier_pdf(f, a.nmi)
        if b.month in bills:
            raise BillingError(f"More than one supplier PDF for {mlabel(b.month)}")
        bills[b.month] = b
        info(f"Supplier PDF {b.file}: {b.account}, {b.start:%d/%m/%Y} to {b.finish:%d/%m/%Y}, "
             f"total {money(b.total)} (NMI {b.nmi})")
    if len({b.nmi for b in bills.values()}) > 1:
        note("Supplier PDFs refer to different NMIs.")
    missing = [m for m in mkeys if m not in bills]
    if missing:
        raise BillingError("No supplier PDF for: " + ", ".join(mlabel(m) for m in missing))

    results = [calc_month(m, hh, bills[m], rd, a) for m in mkeys]
    for r in results:
        c = r.cluster
        ex = ", ".join(f"{x}={c[f'extra_{x}_kwh']:,.0f} kWh" for x in r.extras) or "none"
        info(f"{mlabel(r.month)}: supplier {money(c['supplier_total'])} | households {money(c['final_total'])} | "
             f"variance {money(c['variance'])} | extra meters: {ex}")
        if abs(c["final_total"] - c["supplier_total"]) > 0.005:
            note(f"{mlabel(r.month)}: household totals do not reconcile to the supplier bill!")

    od = a.outdir
    os.makedirs(od, exist_ok=True)
    code = mkeys[0].replace("-", "") + ("" if len(mkeys) == 1 else "-" + mkeys[-1][5:7])
    plabel = mshort(mkeys[0]) + ("" if len(mkeys) == 1 else "-" + mshort(mkeys[-1]))
    accts = [bills[m].account for m in mkeys]
    consecutive = mkeys == [add_months(mkeys[0], k) for k in range(len(mkeys))]
    months_txt = (", ".join(mshort(m) for m in mkeys) if len(mkeys) <= 4 else
                  f"{mshort(mkeys[0])} - {mshort(mkeys[-1])} ({len(mkeys)} months)" if consecutive else
                  f"{len(mkeys)} months")
    ctx = dict(title=a.title, issuer=a.issuer, issue=issue, due=issue + dt.timedelta(days=a.due_days),
               period_code=code, months_txt=months_txt,
               supplier_short=accts[0] if len(accts) == 1 else f"{accts[0]} to {accts[-1]}")

    tabs = os.path.join(od, f"monthly_tabs_{a.year}_{plabel}.xlsx")
    write_monthly_tabs(tabs, results, hh, a)
    info(f"Wrote monthly tabs to {tabs}")

    if not a.no_xero:
        p_inv, p_cn, n_inv, n_cn = write_xero_csvs(od, results, hh, a, ctx)
        info(f"Wrote Xero files: {p_inv} ({n_inv} lines), {p_cn} ({n_cn} credit notes)")

    totals = []
    for h in hh.to_dict("records"):
        fn = os.path.join(od, f"{h['Cluster']}_{h['Lot'].replace(' ', '')}_{a.year}_{plabel}.pdf")
        totals.append(dict(Identifier=h["Identifier"], Owners=h["Owners"],
                           Total=render_invoice(fn, h["Lot"], results, a, h, ctx), File=os.path.basename(fn)))
    info(f"Wrote {len(totals)} invoices to {os.path.abspath(od)}")

    pd.DataFrame(totals).to_csv(os.path.join(od, f"invoice_totals_{a.year}_{plabel}.csv"), index=False)
    pd.DataFrame([r.cluster for r in results]).to_csv(os.path.join(od, f"cluster_monthly_summary_{a.year}_{plabel}.csv"),
                                                      index=False)
    pd.concat([r.bill.lines.assign(month=r.month, account=r.bill.account) for r in results]).to_csv(
        os.path.join(od, f"supplier_bill_lines_{a.year}_{plabel}.csv"), index=False)
    ledger = os.path.join(od, "meter_readings_ledger.csv")
    rd_df.sort_values(["asset", "month"])[LEDGER_COLS].to_csv(ledger, index=False)
    info(f"Wrote {ledger} (use it as --opening-readings next time)")
    if NOTES:
        with open(os.path.join(od, "run_notes.txt"), "w") as fh:
            fh.write("\n".join(NOTES) + "\n")
    info(f"Period total: households {money(sum(t['Total'] for t in totals))} vs supplier "
         f"{money(sum(bills[m].total for m in mkeys))}.")
    return results


def main():
    try:
        run()
    except BillingError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
